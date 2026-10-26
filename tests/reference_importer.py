"""Read the supplied workbook without executing formulas, links, or macros."""
import datetime as dt
import io
import re
import zipfile
import hashlib
import json
import openpyxl

SHEETS = ('砍樹', '圍籬', '雜草垃圾')

def text(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def date(v):
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime('%Y-%m-%d')
    s = text(v)
    if not s: return ''
    try:
        parts = re.split(r'[-/.]', s[:10])
        y, m, d = map(int, parts)
        if y < 1911: y += 1911
        return dt.date(y, m, d).isoformat()
    except (ValueError, TypeError): return ''

def lots(values):
    out = []
    for value in values:
        s = text(value).replace('－', '-').replace('之', '-').replace('號', '')
        for lot in re.split(r'[、,，;；\s]+', s):
            if lot and re.fullmatch(r'\d{1,8}(?:-\d{1,8})?', lot) and lot not in out: out.append(lot)
    return out

def parcels_from(values, district, section):
    """A few rows contain an explicit second township/section in a lot cell."""
    parcels, invalid = [], []
    for value in values:
        s = text(value).replace('－', '-').replace('之', '-').replace('號', '')
        for part in re.split(r'[、,，;；\s]+', s):
            if not part: continue
            if re.fullmatch(r'\d{1,8}(?:-\d{1,8})?', part):
                item={'district':district,'section':section,'lot':part}
            else:
                match=re.fullmatch(r'(?:(?P<district>[\u4e00-\u9fff]{1,6}[鄉鎮市區]))?(?P<section>[\u4e00-\u9fff\d]+段)(?P<lot>\d{1,8}(?:-\d{1,8})?)',part)
                if not match: invalid.append(part); continue
                item={'district':match.group('district') or district,'section':match.group('section'),'lot':match.group('lot')}
            if item not in parcels:parcels.append(item)
    return parcels, invalid

def parse_workbook(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(f.file_size for f in z.infolist()) > 80 * 1024 * 1024 or len(z.infolist()) > 4000:
                raise ValueError('解壓後檔案過大，請分批匯入。')
        w = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError) as e: raise ValueError('無法讀取此 XLSX 檔案。') from e
    records, warnings, seen = [], [], set()
    for name in SHEETS:
        if name not in w.sheetnames: continue
        s = w[name]
        if s.max_row > 30000: raise ValueError('單一工作表超過 30,000 列，請分批匯入。')
        for rownum, row in enumerate(s.values, 1):
            if rownum == 1: continue
            r = list(row) + [None] * 50
            section_idx = 10 if name == '雜草垃圾' else 8
            section = text(r[section_idx])
            if not section: continue  # prefilled serial numbers are not tasks
            district = text(r[3])
            if district.startswith('#'): district = ''
            parcels, invalid = parcels_from(r[section_idx + 1:section_idx + 6],district,section)
            if invalid:
                warnings.append(f'{name} 第 {rownum} 列：地號內容「{"、".join(invalid)}」無法辨識，請修正後再匯入；整列暫不匯入。')
                continue
            if not parcels:
                warnings.append(f'{name} 第 {rownum} 列：無有效地號，未匯入。'); continue
            serial = text(r[0])
            if not serial:
                identity=[text(v) for v in row]
                while identity and not identity[-1]:identity.pop()
                digest=hashlib.sha256(json.dumps(identity,ensure_ascii=False).encode()).hexdigest()[:20]
                base=f'{name}:未編號:{digest}'
                key=base;occurrence=1
                while key in seen:
                    occurrence+=1;key=f'{base}:{occurrence}'
                warnings.append(f'{name} 第 {rownum} 列：缺少序號，已使用內容識別碼保留為一筆任務；建議補齊序號。')
                serial='未編號'
            else:key = f'{name}:{serial}'
            if key in seen:
                warnings.append(f'{name} 第 {rownum} 列：序號 {serial} 重複，未匯入。'); continue
            seen.add(key)
            if not district: warnings.append(f'{name} 第 {rownum} 列：鄉鎮未取得公式快取，請核對。')
            notice_idx = 7 if name == '雜草垃圾' else 5
            notice, due = date(r[notice_idx]), date(r[notice_idx + 1])
            if not notice: warnings.append(f'{name} 第 {rownum} 列：通知日期空白或無法辨識，請設定成果日期。')
            person = text(r[5] if name == '雜草垃圾' else r[29])
            vendor = text(r[6] if name == '雜草垃圾' else r[30])
            items = []
            if name == '雜草垃圾':
                for i in range(3):
                    if text(r[16+i*2]): items.append({'name':text(r[16+i*2]), 'quantity':text(r[17+i*2]), 'unit':'㎡', 'note':text(r[22+i])})
            elif name == '圍籬':
                for i in range(3):
                    if text(r[14+i*2]) or text(r[15+i*2]):
                        items.append({'name':text(r[14+i*2]) or '圍籬', 'quantity':text(r[15+i*2]), 'unit':text(r[20+i]) or '公尺', 'note':text(r[26+i])})
            else:
                for i in range(3):
                    if text(r[15+i*2]): items.append({'name':text(r[23+i]) or '樹木處理', 'quantity':text(r[15+i*2]), 'unit':text(r[20+i]) or '株', 'note':text(r[26+i])})
            records.append({'key':key, 'sheet':name, 'row':rownum, 'serial':serial,
                'case_no': '-'.join(filter(None,[text(r[1]),text(r[2])])), 'district':district,
                'section_code':text(r[4]), 'section':section, 'lots':[p['lot'] for p in parcels], 'parcels':parcels, 'notice':notice, 'due':due,
                'area':text(r[9] if name == '雜草垃圾' else r[7]), 'source_person':person,
                'vendor':vendor, 'items':items})
    w.close()
    if not records: raise ValueError('找不到可匯入資料。請使用含「砍樹／圍籬／雜草垃圾」工作表的工作簿。')
    return {'records':records, 'warnings':warnings,
        'summary':{s:{'groups':sum(r['sheet']==s for r in records),'tasks':sum(r['sheet']==s for r in records),'parcels':sum(len(r['lots']) for r in records if r['sheet']==s)} for s in SHEETS}}
